"""
Reports 舊版流程的狀態修正測試（LAUNCH_PLAN P1）

修正前的兩個問題：
1. FormFillService._templates 是 instance attribute，而每個 endpoint
   每請求 new 一個 service → 模板活不過單一請求，/reports 流程跨請求
   必然 "Template not found"
2. /generate 為背景任務 + status 輪詢，狀態存行程記憶體 →
   多實例/scale-to-zero 下輪詢 404

修正後：
- _templates 為 class-level（行程內跨請求共享）
- /generate 同步執行，直接回 completed + download_url
- Template not found → 404（原為 500）
"""

import io
import os
import sys

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
os.environ.setdefault("GEMINI_API_KEY", "test-key")

from app.main import app  # noqa: E402
from app.services.form_fill import FormFillService  # noqa: E402

TEMPLATE_ID = "test-report-template-001"


def _make_xlsx_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "設備名稱"
    ws["A2"] = "檢查結果"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def seeded_template():
    """塞入一個最小可用模板（generate_report 消費 file_type/file_content/fields）"""
    FormFillService._templates[TEMPLATE_ID] = {
        "id": TEMPLATE_ID,
        "name": "測試模板",
        "vendor_name": "測試廠商",
        "file_type": "xlsx",
        "file_content": _make_xlsx_template_bytes(),
        "fields": [
            {"name": "設備名稱", "location": "B1", "mapping": "equipment_name"},
            {"name": "檢查結果", "location": "B2", "mapping": "condition_assessment"},
        ],
    }
    yield TEMPLATE_ID
    FormFillService._templates.pop(TEMPLATE_ID, None)


def _inspection_data() -> dict:
    return {
        "inspection_id": "INSP-TEST-001",
        "equipment_name": "測試馬達",
        "equipment_type": "馬達",
        "inspection_date": "2026-08-31",
        "condition_assessment": "狀況良好",
    }


def test_templates_shared_across_service_instances(seeded_template):
    """class-level 存放區：instance A 塞入的模板，instance B 看得到。"""
    another_instance = FormFillService()
    assert TEMPLATE_ID in another_instance._templates


def test_generate_is_synchronous_and_downloadable(client, seeded_template):
    """POST /generate 同步完成 → 立即回 completed + download_url → 可下載。"""
    resp = client.post("/api/reports/generate", json={
        "template_id": TEMPLATE_ID,
        "inspection_data": _inspection_data(),
        "output_format": "xlsx",
    })
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["success"] is True
    assert data["status"] == "completed"
    assert data["download_url"], "同步完成必須直接給 download_url"

    report_id = data["report_id"]

    # status 立即可查（同一行程）
    status = client.get(f"/api/reports/{report_id}/status")
    assert status.status_code == 200
    assert status.json()["status"] == "completed"

    # 檔案立即可下載且為 xlsx
    download = client.get(f"/api/reports/{report_id}/download")
    assert download.status_code == 200
    assert download.content[:2] == b"PK", "xlsx 應為 zip 容器（PK 開頭）"


def test_generate_unknown_template_returns_404(client):
    resp = client.post("/api/reports/generate", json={
        "template_id": "no-such-template",
        "inspection_data": _inspection_data(),
        "output_format": "xlsx",
    })
    assert resp.status_code == 404


def test_preview_unknown_template_returns_404(client):
    resp = client.post("/api/reports/preview", json={
        "template_id": "no-such-template",
        "inspection_data": _inspection_data(),
    })
    assert resp.status_code == 404


def test_unknown_report_status_returns_404(client):
    resp = client.get("/api/reports/00000000-0000-0000-0000-000000000000/status")
    assert resp.status_code == 404
